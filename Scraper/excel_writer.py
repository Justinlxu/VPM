"""
Excel Writer

Writes scraped match data to a 3-sheet Excel workbook.
Sheet 1: Model Features (one row per map)
Sheet 2: Raw Round Data (one row per map, round-by-round)
Sheet 3: Player Stats (10 rows per map)
"""

import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import (
    OUTPUT_FILE, MAX_ROUNDS,
    ECO_LOADOUT_DIFF, GUN_ROUND_LOADOUT,
    HALF_STARTS, BONUS_STREAK_LENGTH,
)


# ══════════════════════════════════════════════════════════════════
# COLUMN DEFINITIONS
# ══════════════════════════════════════════════════════════════════

SHEET1_HEADERS = [
    "Match ID", "Date", "Map",
    "Team A", "Team B", "Winner",
    "Score A", "Score B",
    "FK/FD Diff A", "FK/FD Diff B",
    "Pistols Won A", "Pistols Won B",
    "ACS A", "ACS B",
    "KAST A", "KAST B",
    "Bonus Won A", "Bonus Won B",
    "Eco Won A", "Eco Won B",
    "Gun Rounds Won A", "Gun Rounds Won B",
]

def build_sheet2_headers():
    headers = ["Match ID"]
    for r in range(1, MAX_ROUNDS + 1):
        headers.append(f"R{r} Winner")
        headers.append(f"R{r} Loadout A")
        headers.append(f"R{r} Loadout B")
    return headers

SHEET2_HEADERS = build_sheet2_headers()

SHEET3_HEADERS = [
    "Match ID", "Player Name", "Team", "Agent",
    "ACS", "Kills", "Deaths", "Assists",
    "FK", "FD", "KAST",
]


# ══════════════════════════════════════════════════════════════════
# STYLING
# ══════════════════════════════════════════════════════════════════

HEADER_FONT = Font(name="Arial", bold=True, size=10, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F3136")  # Dark background
DATA_FONT = Font(name="Arial", size=10)
BORDER = Border(
    bottom=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
)

TEAM_A_FILL = PatternFill("solid", fgColor="E8F5E9")  # Light green
TEAM_B_FILL = PatternFill("solid", fgColor="E3F2FD")  # Light blue
WINNER_FILL = PatternFill("solid", fgColor="FFF9C4")  # Light yellow


# ══════════════════════════════════════════════════════════════════
# WORKBOOK CREATION
# ══════════════════════════════════════════════════════════════════

def create_workbook():
    """Create a new Excel workbook with 3 sheets and headers."""
    wb = Workbook()

    # Sheet 1: Model Features
    ws1 = wb.active
    ws1.title = "Model Features"
    _write_headers(ws1, SHEET1_HEADERS)
    _set_column_widths(ws1, SHEET1_HEADERS, default_width=14)

    # Sheet 2: Round Data
    ws2 = wb.create_sheet("Round Data")
    _write_headers(ws2, SHEET2_HEADERS)
    _set_column_widths(ws2, SHEET2_HEADERS, default_width=10)

    # Sheet 3: Player Stats
    ws3 = wb.create_sheet("Player Stats")
    _write_headers(ws3, SHEET3_HEADERS)
    _set_column_widths(ws3, SHEET3_HEADERS, default_width=14)

    wb.save(OUTPUT_FILE)
    print(f"Created {OUTPUT_FILE} with 3 sheets")
    return wb


def _write_headers(ws, headers):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _set_column_widths(ws, headers, default_width=14):
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        width = max(len(header) + 2, default_width)
        # Cap width for round data columns
        if "Loadout" in header or "R" in header[:2]:
            width = 12
        ws.column_dimensions[col_letter].width = min(width, 20)


# ══════════════════════════════════════════════════════════════════
# WRITING DATA
# ══════════════════════════════════════════════════════════════════

def load_or_create_workbook():
    """Load existing workbook or create a new one."""
    if os.path.exists(OUTPUT_FILE):
        wb = load_workbook(OUTPUT_FILE)
        print(f"Loaded existing {OUTPUT_FILE}")
    else:
        wb = create_workbook()
    return wb


def write_match_data(wb, match_data):
    """
    Write a single map's data to all 3 sheets.

    Args:
        wb: openpyxl Workbook
        match_data: dict from scraper.scrape_match() for one map
    """
    _write_sheet1_row(wb["Model Features"], match_data)
    _write_sheet2_row(wb["Round Data"], match_data)
    _write_sheet3_rows(wb["Player Stats"], match_data)


def _write_sheet1_row(ws, data):
    """Write one row to Sheet 1 (Model Features)."""
    next_row = ws.max_row + 1

    row = [
        data.get("match_id", ""),
        data.get("date", ""),
        data.get("map", ""),
        data.get("team_a", ""),
        data.get("team_b", ""),
        data.get("winner", ""),
        data.get("score_a", 0),
        data.get("score_b", 0),
        data.get("fk_fd_diff_a", 0),
        data.get("fk_fd_diff_b", 0),
        data.get("pistols_won_a", 0),
        data.get("pistols_won_b", 0),
        round(data.get("acs_a", 0), 1),
        round(data.get("acs_b", 0), 1),
        round(data.get("kast_a", 0), 1),
        round(data.get("kast_b", 0), 1),
        data.get("bonus_won_a", 0),
        data.get("bonus_won_b", 0),
        data.get("eco_won_a", 0),
        data.get("eco_won_b", 0),
        data.get("gun_won_a", 0),
        data.get("gun_won_b", 0),
    ]

    for col_idx, value in enumerate(row, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.font = DATA_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    # Highlight the winning team's score column
    winner = data.get("winner", "")
    if winner == data.get("team_a", ""):
        winner_col = 7  # Score A
    elif winner == data.get("team_b", ""):
        winner_col = 8  # Score B
    else:
        winner_col = None
    if winner_col:
        ws.cell(row=next_row, column=winner_col).fill = WINNER_FILL


def _write_sheet2_row(ws, data):
    """Write one row to Sheet 2 (Round Data)."""
    next_row = ws.max_row + 1
    rounds = data.get("rounds", [])

    # Build the round data
    rmap = {r["round"]: r for r in rounds}
    row = [data.get("match_id", "")]

    for r_num in range(1, MAX_ROUNDS + 1):
        r = rmap.get(r_num)
        if r:
            winner_str = data.get("team_a", "A") if r["winner"] == "a" else (
                data.get("team_b", "B") if r["winner"] == "b" else ""
            )
            row.append(winner_str)
            row.append(r.get("loadout_a"))
            row.append(r.get("loadout_b"))
        else:
            row.append("")   # Winner
            row.append(None) # Loadout A
            row.append(None) # Loadout B

    for col_idx, value in enumerate(row, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.font = DATA_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

        # Color round winner cells
        if col_idx > 1 and (col_idx - 2) % 3 == 0:  # Winner columns
            if value == data.get("team_a", ""):
                cell.fill = TEAM_A_FILL
            elif value == data.get("team_b", ""):
                cell.fill = TEAM_B_FILL


def _write_sheet3_rows(ws, data):
    """Write player rows to Sheet 3 (10 rows per map — 5 per team)."""
    match_id = data.get("match_id", "")
    team_a = data.get("team_a", "")
    team_b = data.get("team_b", "")

    for player in data.get("players_a", []):
        _write_player_row(ws, match_id, player, team_a)

    for player in data.get("players_b", []):
        _write_player_row(ws, match_id, player, team_b)


def _write_player_row(ws, match_id, player, team_name):
    """Write a single player stat row to Sheet 3."""
    next_row = ws.max_row + 1

    row = [
        match_id,
        player.get("name", ""),
        team_name,
        player.get("agent", ""),
        player.get("acs", 0),
        player.get("kills", 0),
        player.get("deaths", 0),
        player.get("assists", 0),
        player.get("fk", 0),
        player.get("fd", 0),
        player.get("kast", 0),
    ]

    for col_idx, value in enumerate(row, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.font = DATA_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")


def save_workbook(wb):
    """Save the workbook to disk."""
    wb.save(OUTPUT_FILE)


# ══════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════

def print_summary(wb):
    """Print a summary of what's in the workbook."""
    ws1 = wb["Model Features"]
    ws3 = wb["Player Stats"]

    n_maps = ws1.max_row - 1
    n_players = ws3.max_row - 1

    print(f"\n{'═' * 40}")
    print(f"  WORKBOOK SUMMARY")
    print(f"{'═' * 40}")
    print(f"  Sheet 1 (Model Features): {n_maps} maps")
    print(f"  Sheet 2 (Round Data):     {n_maps} maps")
    print(f"  Sheet 3 (Player Stats):   {n_players} player rows")
    print(f"  Saved to: {OUTPUT_FILE}")


# ══════════════════════════════════════════════════════════════════
# POST-PROCESSING: DERIVE ECONOMY FEATURES FROM ROUND DATA
# ══════════════════════════════════════════════════════════════════

def derive_features_from_rounds(wb):
    """
    Read the Round Data sheet and (re)compute Bonus Won, Eco Won, and
    Gun Rounds Won in the Model Features sheet.

    Definitions (from config.py):
      Bonus    — team wins all BONUS_STREAK_LENGTH rounds from each half start
                 (rounds 1-3 and 13-15 by default)
      Eco      — |loadout_A - loadout_B| >= ECO_LOADOUT_DIFF (12 500);
                 eco_won_X counts rounds the ECO side (lower loadout) won
      Gun      — both loadouts >= GUN_ROUND_LOADOUT (16 000);
                 gun_won_X counts those rounds won by team X

    Safe to re-run: always overwrites with freshly computed values.
    Returns the number of map rows updated.
    """
    ws_feat = wb["Model Features"]
    ws_rnd  = wb["Round Data"]

    # Model Features column positions (1-based, matching SHEET1_HEADERS)
    BONUS_A_COL, BONUS_B_COL = 17, 18
    ECO_A_COL,   ECO_B_COL   = 19, 20
    GUN_A_COL,   GUN_B_COL   = 21, 22

    # Build a Match-ID → row-number index for Model Features
    feat_index = {}
    for row_idx in range(2, ws_feat.max_row + 1):
        mid = ws_feat.cell(row=row_idx, column=1).value
        if mid is not None:
            feat_index[str(mid)] = row_idx

    updated = 0
    for rnd_row in range(2, ws_rnd.max_row + 1):
        mid = ws_rnd.cell(row=rnd_row, column=1).value
        if mid is None:
            continue

        feat_row = feat_index.get(str(mid))
        if feat_row is None:
            continue

        # Team names from Model Features (needed to convert winner name → a/b)
        team_a = str(ws_feat.cell(row=feat_row, column=4).value or "")
        team_b = str(ws_feat.cell(row=feat_row, column=5).value or "")

        # Round Data column layout: Match ID | R1 Winner | R1 Loadout A | R1 Loadout B | R2 …
        # For round r: base_col = 2 + (r-1)*3
        rounds = {}
        for r in range(1, MAX_ROUNDS + 1):
            base = 2 + (r - 1) * 3
            winner_name = ws_rnd.cell(row=rnd_row, column=base).value
            la = ws_rnd.cell(row=rnd_row, column=base + 1).value
            lb = ws_rnd.cell(row=rnd_row, column=base + 2).value

            if not winner_name:
                continue

            side = ("a" if str(winner_name) == team_a
                    else "b" if str(winner_name) == team_b
                    else "")
            if side:
                rounds[r] = {"winner": side, "la": la, "lb": lb}

        # ── Bonus ──────────────────────────────────────────────────
        bonus_a = bonus_b = 0
        for half_start in HALF_STARTS:
            winners = [rounds.get(half_start + i, {}).get("winner")
                       for i in range(BONUS_STREAK_LENGTH)]
            if all(w == "a" for w in winners):
                bonus_a += 1
            if all(w == "b" for w in winners):
                bonus_b += 1

        # ── Eco and Gun ────────────────────────────────────────────
        eco_a = eco_b = gun_a = gun_b = 0
        for r_data in rounds.values():
            la, lb, w = r_data["la"], r_data["lb"], r_data["winner"]
            if la is None or lb is None:
                continue
            if la >= GUN_ROUND_LOADOUT and lb >= GUN_ROUND_LOADOUT:
                if w == "a": gun_a += 1
                elif w == "b": gun_b += 1
            elif abs(la - lb) >= ECO_LOADOUT_DIFF:
                if la < lb and w == "a": eco_a += 1
                elif lb < la and w == "b": eco_b += 1

        # ── Write back ─────────────────────────────────────────────
        for col, val in [
            (BONUS_A_COL, bonus_a), (BONUS_B_COL, bonus_b),
            (ECO_A_COL,   eco_a),   (ECO_B_COL,   eco_b),
            (GUN_A_COL,   gun_a),   (GUN_B_COL,   gun_b),
        ]:
            cell = ws_feat.cell(row=feat_row, column=col, value=val)
            cell.font = DATA_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")

        updated += 1

    print(f"  Derived economy features for {updated} maps")
    return updated
